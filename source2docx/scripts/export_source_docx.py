#!/usr/bin/env python3
"""
Export project source code into .docx for software copyright submission.

Strategy: flatten source -> paginate in memory -> write master.docx ->
split by page ranges into source-doc-NN.docx (same layout, guaranteed page counts).

Requires: pip install python-docx
"""

from __future__ import annotations

import argparse
import fnmatch
import json
import math
import re
import sys
from collections import deque
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

try:
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Mm, Pt, Twips
except ImportError:
    print("Error: python-docx is required. Install with: pip install python-docx", file=sys.stderr)
    sys.exit(1)

# 五号 = 10.5pt；Word「宋体（正文）」
DEFAULT_FONT_NAME = "宋体"
DEFAULT_FONT_SIZE_PT = 10.5

DEFAULT_SKIP_DIRS = {
    "node_modules",
    ".git",
    "dist",
    "build",
    "coverage",
    ".next",
    "vendor",
    "__pycache__",
    ".venv",
    "venv",
}

DEFAULT_SKIP_GLOBS = [
    "**/package-lock.json",
    "**/yarn.lock",
    "**/pnpm-lock.yaml",
    "**/*.min.js",
    "**/*.map",
    "**/.env",
    "**/.env.*",
    "**/*.pem",
    "**/*.key",
    "**/*.p12",
    "**/*.png",
    "**/*.jpg",
    "**/*.jpeg",
    "**/*.gif",
    "**/*.svg",
    "**/*.ico",
    "**/*.woff",
    "**/*.woff2",
    "**/*.ttf",
    "**/*.mp4",
    "**/*.zip",
]

SOURCE_EXTENSIONS = {
    ".js", ".jsx", ".ts", ".tsx", ".vue", ".py", ".go", ".java", ".kt", ".rs",
    ".c", ".cpp", ".h", ".hpp", ".cs", ".rb", ".php", ".swift", ".m", ".mm",
    ".sql", ".scss", ".less", ".css", ".html", ".xml", ".json", ".yaml", ".yml", ".md", ".sh",
}

IMPORT_PATTERNS = [
    re.compile(r"""import\s+(?:[\w*{}\s,]+\s+from\s+)?['"]([^'"]+)['"]"""),
    re.compile(r"""require\s*\(\s*['"]([^'"]+)['"]\s*\)"""),
    re.compile(r"""import\s*\(\s*['"]([^'"]+)['"]\s*\)"""),
    re.compile(r"""from\s+['"]([^'"]+)['"]\s+import"""),
    re.compile(r"""^import\s+([a-zA-Z0-9_.]+)""", re.MULTILINE),
    re.compile(r"""^from\s+([a-zA-Z0-9_.]+)\s+import""", re.MULTILINE),
    re.compile(r"""#include\s+[<"]([^>"]+)[>"]"""),
]


@dataclass
class PageLayout:
    page_width: Mm | Inches
    page_height: Mm | Inches
    margin_top: Cm
    margin_bottom: Cm
    margin_left: Cm
    margin_right: Cm
    lines_per_page: int
    font_name: str = DEFAULT_FONT_NAME
    font_size_pt: float = DEFAULT_FONT_SIZE_PT

    @property
    def content_height_twips(self) -> int:
        return (
            int(self.page_height.twips)
            - int(self.margin_top.twips)
            - int(self.margin_bottom.twips)
        )

    @property
    def line_spacing_twips(self) -> int:
        # 整 twips 向下取整，保证 lines_per_page 行不会超出可排版高度
        return self.content_height_twips // self.lines_per_page

    @property
    def line_spacing_pt(self) -> float:
        return self.line_spacing_twips / 20

    def as_dict(self) -> dict:
        return {
            "page_width_mm": round(self.page_width.mm, 2),
            "page_height_mm": round(self.page_height.mm, 2),
            "margin_top_cm": round(self.margin_top.cm, 2),
            "margin_bottom_cm": round(self.margin_bottom.cm, 2),
            "margin_left_cm": round(self.margin_left.cm, 2),
            "margin_right_cm": round(self.margin_right.cm, 2),
            "lines_per_page": self.lines_per_page,
            "line_spacing_twips": self.line_spacing_twips,
            "line_spacing_pt": round(self.line_spacing_pt, 3),
            "font_name": self.font_name,
            "font_size_pt": self.font_size_pt,
        }


@dataclass
class FileSegment:
    relative_path: str
    lines: list[str]
    obfuscated: bool = False


def build_page_layout(
    page_size: str,
    lines_per_page: int,
    margin_top_cm: float,
    margin_bottom_cm: float,
    margin_left_cm: float,
    margin_right_cm: float,
    font_name: str,
    font_size_pt: float,
) -> PageLayout:
    if page_size == "letter":
        page_width, page_height = Inches(8.5), Inches(11)
    else:
        page_width, page_height = Mm(210), Mm(297)

    return PageLayout(
        page_width=page_width,
        page_height=page_height,
        margin_top=Cm(margin_top_cm),
        margin_bottom=Cm(margin_bottom_cm),
        margin_left=Cm(margin_left_cm),
        margin_right=Cm(margin_right_cm),
        lines_per_page=lines_per_page,
        font_name=font_name,
        font_size_pt=font_size_pt,
    )


def apply_page_layout(section, layout: PageLayout) -> None:
    section.page_width = layout.page_width
    section.page_height = layout.page_height
    section.top_margin = layout.margin_top
    section.bottom_margin = layout.margin_bottom
    section.left_margin = layout.margin_left
    section.right_margin = layout.margin_right
    remove_doc_grid(section)


def remove_doc_grid(section) -> None:
    """去掉文档网格，避免与段落行距冲突产生整页空白。"""
    sect_pr = section._sectPr
    for old in sect_pr.findall(qn("w:docGrid")):
        sect_pr.remove(old)


def clear_body_paragraphs(doc: Document) -> None:
    """移除 Document() 默认空段落，避免首页多占一行导致分页错位。"""
    body = doc.element.body
    for child in list(body):
        if child.tag == qn("w:p"):
            body.remove(child)


def configure_normal_style(doc: Document, layout: PageLayout) -> None:
    style = doc.styles["Normal"]
    style.font.name = layout.font_name
    style.font.size = Pt(layout.font_size_pt)
    fmt = style.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Twips(layout.line_spacing_twips)


def configure_code_paragraph(p, layout: PageLayout) -> None:
    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Twips(layout.line_spacing_twips)
    fmt.widow_control = False
    fmt.keep_together = False
    for run in p.runs:
        set_run_font(run, layout)


def set_run_font(run, layout: PageLayout) -> None:
    run.font.name = layout.font_name
    run.font.size = Pt(layout.font_size_pt)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for tag in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{tag}"), layout.font_name)


def sanitize_line(line: str) -> str:
    """去掉换页/换行控制符，保留缩进空格。"""
    text = line.replace("\r", "").replace("\n", "")
    for ch in ("\f", "\v", "\x0b"):
        text = text.replace(ch, "")
    return text


def add_line_paragraph(doc: Document, raw: str, layout: PageLayout) -> None:
    """一行源码 = 一个段落。空行不用空格占位（空格会导致 Word 整页空白）。"""
    text = sanitize_line(raw)
    p = doc.add_paragraph()

    if text:
        run = p.add_run(text)
    else:
        run = p.add_run("")
        t_elem = run._element.find(qn("w:t"))
        if t_elem is not None:
            t_elem.set(qn("xml:space"), "preserve")

    configure_code_paragraph(p, layout)


def normalize_rel(path: Path, root: Path) -> str:
    rel = path.relative_to(root).as_posix()
    if rel.startswith("./"):
        rel = rel[2:]
    return rel


def should_skip(rel_posix: str, skip_globs: list[str]) -> bool:
    parts = rel_posix.split("/")
    if any(p in DEFAULT_SKIP_DIRS for p in parts):
        return True
    for pattern in skip_globs:
        if fnmatch.fnmatch(rel_posix, pattern) or fnmatch.fnmatch(rel_posix.split("/")[-1], pattern):
            return True
    return False


def resolve_import(ref: str, source_file: Path, root: Path) -> Path | None:
    if ref.startswith((".", "/")):
        pass
    elif ref.startswith("@/"):
        ref = "src/" + ref[2:]
    else:
        return None

    if ref.startswith("/"):
        base = root / ref.lstrip("/")
    else:
        base = (source_file.parent / ref).resolve()

    try:
        base.relative_to(root.resolve())
    except ValueError:
        return None

    candidates: list[Path] = []
    if base.suffix:
        candidates.append(base)
    else:
        for ext in SOURCE_EXTENSIONS:
            candidates.append(Path(str(base) + ext))
            candidates.append(base / ("index" + ext))

    for c in candidates:
        if c.is_file():
            return c.resolve()
    return None


def parse_imports(content: str, source_file: Path, root: Path) -> list[Path]:
    found: list[Path] = []
    seen: set[str] = set()
    for pat in IMPORT_PATTERNS:
        for m in pat.finditer(content):
            ref = m.group(1)
            if ref.startswith(".") or ref.startswith("/") or ref.startswith("@/"):
                resolved = resolve_import(ref, source_file, root)
                if resolved:
                    key = normalize_rel(resolved, root)
                    if key not in seen:
                        seen.add(key)
                        found.append(resolved)
    return found


def discover_files(entry: Path, root: Path, skip_globs: list[str]) -> list[Path]:
    if not entry.is_file():
        raise FileNotFoundError(f"Entry file not found: {entry}")

    rel_entry = normalize_rel(entry, root)
    if should_skip(rel_entry, skip_globs):
        raise ValueError(f"Entry file is excluded: {rel_entry}")

    queue: deque[Path] = deque([entry.resolve()])
    visited: set[str] = set()
    ordered: list[Path] = []

    while queue:
        current = queue.popleft()
        rel = normalize_rel(current, root)
        if rel in visited or should_skip(rel, skip_globs):
            continue
        if not current.is_file():
            continue
        if current.suffix.lower() not in SOURCE_EXTENSIONS and current.suffix:
            continue

        visited.add(rel)
        ordered.append(current)

        try:
            content = current.read_text(encoding="utf-8", errors="replace")
        except OSError:
            continue

        for dep in parse_imports(content, current, root):
            dep_rel = normalize_rel(dep, root)
            if dep_rel not in visited:
                queue.append(dep)

    return ordered


def _mask_sensitive_functions(lines: list[str], level: str) -> tuple[list[str], bool]:
    if level not in ("medium", "strict"):
        return lines, False

    sensitive = re.compile(
        r"(encrypt|decrypt|sign|verify|hash|hmac|license|drm|billing)",
        re.IGNORECASE,
    )
    result: list[str] = []
    changed = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if sensitive.search(line) and ("function" in line or "=>" in line):
            result.append(line)
            i += 1
            if i < len(lines) and "{" in lines[i] and "}" not in lines[i]:
                result.append(lines[i])
                depth = lines[i].count("{") - lines[i].count("}")
                i += 1
            elif "{" in line:
                depth = line.count("{") - line.count("}")
            else:
                continue
            while i < len(lines) and depth > 0:
                body = lines[i]
                if "{" in body:
                    depth += body.count("{")
                if "}" in body:
                    depth -= body.count("}")
                if depth == 0:
                    result.append(body)
                    i += 1
                    break
                indent = re.match(r"^(\s*)", body).group(1)
                result.append(f"{indent}// core logic omitted for copyright submission")
                changed = True
                i += 1
            continue
        result.append(line)
        i += 1
    return result, changed


def obfuscate_content(content: str, level: str) -> tuple[str, bool]:
    changed = False
    lines = content.splitlines()
    out: list[str] = []

    for line in lines:
        original = line
        line = re.sub(
            r"(?i)(api[_-]?key|secret|password|token|private[_-]?key|access[_-]?key)\s*[:=]\s*['\"][^'\"]+['\"]",
            r'\1 = "***"',
            line,
        )
        line = re.sub(
            r"(?i)(mongodb|redis|mysql|postgres(?:ql)?)://[^\s'\"]+",
            '"<connection-redacted>"',
            line,
        )
        line = re.sub(r"['\"][A-Za-z0-9+/]{40,}={0,2}['\"]", '"<redacted>"', line)
        line = re.sub(
            r"eyJ[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+",
            "<jwt-redacted>",
            line,
        )
        if level == "strict":
            line = re.sub(
                r"^(\s*(?:const|let|var)\s+\w+\s*=\s*)['\"][^'\"]{33,}['\"]",
                r'\1"<redacted>"',
                line,
            )
            line = re.sub(r"/[^/\n]{20,}/[gimsuy]*", "/.../", line)
        if line != original:
            changed = True
        out.append(line)

    out, fn_changed = _mask_sensitive_functions(out, level)
    changed = changed or fn_changed
    return "\n".join(out), changed


def segment_to_lines(seg: FileSegment) -> list[str]:
    return [f"==={seg.relative_path}==="] + seg.lines


def segments_to_lines(segments: Iterable[FileSegment]) -> list[str]:
    lines: list[str] = []
    for seg in segments:
        lines.extend(segment_to_lines(seg))
    return lines


def collect_segments(
    files: Iterable[Path],
    root: Path,
    target_pages: int,
    lines_per_page: int,
    obfuscation: str,
) -> list[FileSegment]:
    """Add whole files until flattened line count supports target_pages."""
    segments: list[FileSegment] = []
    needed_lines = target_pages * lines_per_page

    for fpath in files:
        rel = normalize_rel(fpath, root)
        raw = fpath.read_text(encoding="utf-8", errors="replace")
        obfuscated_text, was_obfuscated = obfuscate_content(raw, obfuscation)
        seg = FileSegment(
            relative_path=rel,
            lines=obfuscated_text.splitlines(),
            obfuscated=was_obfuscated,
        )
        segments.append(seg)
        if len(segments_to_lines(segments)) >= needed_lines:
            break

    if not segments:
        raise ValueError("No source files collected. Check entry file and exclusions.")

    actual_lines = len(segments_to_lines(segments))
    actual_pages = math.ceil(actual_lines / lines_per_page)
    if actual_pages < target_pages:
        raise ValueError(
            f"Source insufficient: need {target_pages} docx pages "
            f"({needed_lines} lines), got {actual_pages} pages ({actual_lines} lines). "
            "Add more entry dependencies or lower total_pages."
        )
    return segments


def paginate_lines(lines: list[str], lines_per_page: int) -> list[list[str]]:
    """Split flat lines into fixed-size pages (docx row units)."""
    pages: list[list[str]] = []
    for i in range(0, len(lines), lines_per_page):
        page = lines[i : i + lines_per_page]
        if page:
            pages.append(page)
    return pages


def truncate_lines_for_pages(lines: list[str], target_pages: int, lines_per_page: int) -> list[str]:
    """Keep exactly target_pages * lines_per_page lines (may end mid-file)."""
    return lines[: target_pages * lines_per_page]


def distribute_page_ranges(total_pages: int, doc_count: int) -> list[tuple[int, int]]:
    """Return [start, end) page index ranges for each output document."""
    if doc_count < 1:
        raise ValueError("doc_count must be >= 1")
    base, rem = divmod(total_pages, doc_count)
    ranges: list[tuple[int, int]] = []
    start = 0
    for i in range(doc_count):
        count = base + (1 if i < rem else 0)
        ranges.append((start, start + count))
        start += count
    return ranges


def flatten_pages(pages: list[list[str]]) -> list[str]:
    lines: list[str] = []
    for page_lines in pages:
        lines.extend(page_lines)
    return lines


def write_docx_from_pages(
    pages: list[list[str]],
    output_path: Path,
    layout: PageLayout,
) -> None:
    """连续段落排版，不插入手动分页符，空行不用空格占位。"""
    lines = flatten_pages(pages)
    if not lines:
        return

    doc = Document()
    apply_page_layout(doc.sections[0], layout)
    configure_normal_style(doc, layout)
    clear_body_paragraphs(doc)

    for raw in lines:
        add_line_paragraph(doc, raw, layout)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def write_manifest(
    path: Path,
    params: dict,
    segments: list[FileSegment],
    page_ranges: list[tuple[int, int]],
    total_pages: int,
) -> None:
    docs = []
    for idx, (start, end) in enumerate(page_ranges, start=1):
        docs.append({
            "file": f"source-doc-{idx:02d}.docx",
            "page_start": start + 1,
            "page_end": end,
            "page_count": end - start,
        })

    payload = {
        "params": params,
        "master": "source-master.docx",
        "total_pages": total_pages,
        "documents": docs,
        "source_segments": [
            {
                "path": s.relative_path,
                "source_lines": len(s.lines),
                "obfuscated": s.obfuscated,
            }
            for s in segments
        ],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Export source code to docx for software copyright.")
    parser.add_argument("--project-root", type=Path, default=Path("."))
    parser.add_argument("--entry", required=True, help="Entry file relative to project root")
    parser.add_argument("--doc-count", type=int, required=True)
    parser.add_argument("--lines-per-page", type=int, required=True)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--total-lines", type=int)
    group.add_argument("--total-pages", type=int)
    parser.add_argument("--obfuscation", choices=["light", "medium", "strict"], default="medium")
    parser.add_argument("--output-dir", type=Path, default=Path("soft-copyright-output"))
    parser.add_argument("--exclude", action="append", default=[], help="Additional glob patterns")
    parser.add_argument("--page-size", choices=["a4", "letter"], default="a4")
    parser.add_argument("--margin-cm", type=float, help="Uniform margin (cm) for all sides")
    parser.add_argument("--margin-top-cm", type=float, default=2.5)
    parser.add_argument("--margin-bottom-cm", type=float, default=2.5)
    parser.add_argument("--margin-left-cm", type=float, default=2.0)
    parser.add_argument("--margin-right-cm", type=float, default=2.0)
    parser.add_argument("--font-name", default=DEFAULT_FONT_NAME)
    parser.add_argument("--font-size", type=float, default=DEFAULT_FONT_SIZE_PT, help="Font size in pt (五号=10.5)")
    parser.add_argument("--skip-master", action="store_true", help="Do not write source-master.docx")
    args = parser.parse_args()

    if args.margin_cm is not None:
        args.margin_top_cm = args.margin_bottom_cm = args.margin_left_cm = args.margin_right_cm = args.margin_cm

    layout = build_page_layout(
        page_size=args.page_size,
        lines_per_page=args.lines_per_page,
        margin_top_cm=args.margin_top_cm,
        margin_bottom_cm=args.margin_bottom_cm,
        margin_left_cm=args.margin_left_cm,
        margin_right_cm=args.margin_right_cm,
        font_name=args.font_name,
        font_size_pt=args.font_size,
    )

    root = args.project_root.resolve()
    entry = (root / args.entry).resolve()
    skip_globs = DEFAULT_SKIP_GLOBS + args.exclude

    target_pages = args.total_pages or math.ceil(args.total_lines / args.lines_per_page)

    discovered = discover_files(entry, root, skip_globs)
    segments = collect_segments(discovered, root, target_pages, args.lines_per_page, args.obfuscation)

    flat_lines = segments_to_lines(segments)
    flat_lines = truncate_lines_for_pages(flat_lines, target_pages, args.lines_per_page)
    all_pages = paginate_lines(flat_lines, args.lines_per_page)

    if len(all_pages) != target_pages:
        raise RuntimeError(
            f"Internal pagination error: expected {target_pages} pages, got {len(all_pages)}"
        )

    out_dir = (root / args.output_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)

    if not args.skip_master:
        write_docx_from_pages(all_pages, out_dir / "source-master.docx", layout)

    page_ranges = distribute_page_ranges(target_pages, args.doc_count)
    for idx, (start, end) in enumerate(page_ranges, start=1):
        chunk = all_pages[start:end]
        if not chunk:
            continue
        write_docx_from_pages(chunk, out_dir / f"source-doc-{idx:02d}.docx", layout)

    params = {
        "entry": args.entry,
        "doc_count": args.doc_count,
        "lines_per_page": args.lines_per_page,
        "target_pages": target_pages,
        "target_docx_lines": target_pages * args.lines_per_page,
        "obfuscation": args.obfuscation,
        "output_dir": str(args.output_dir),
        "layout": layout.as_dict(),
    }
    write_manifest(out_dir / "manifest.json", params, segments, page_ranges, target_pages)

    split_summary = ", ".join(
        f"doc{i:02d}={end - start}p" for i, (start, end) in enumerate(page_ranges, start=1)
    )
    print(
        f"Done -> {out_dir}\n"
        f"Master: {target_pages} pages x {args.lines_per_page} lines/page "
        f"({target_pages * args.lines_per_page} docx lines)\n"
        f"Split: {split_summary}\n"
        f"Font: {layout.font_name} {layout.font_size_pt}pt, "
        f"line spacing {layout.line_spacing_pt:.2f}pt ({layout.line_spacing_twips} twips)"
    )


if __name__ == "__main__":
    main()
